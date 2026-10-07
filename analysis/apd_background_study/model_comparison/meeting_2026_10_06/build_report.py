"""Build PDF and editable DOCX from verified saved-fit comparison data."""
from pathlib import Path
import json,re,html
import numpy as np
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image,Table,TableStyle,PageBreak,KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.enums import TA_LEFT
from docx import Document
from docx.shared import Inches,Pt,RGBColor
R=Path(__file__).resolve().parent;OUT=R/'output';OUT.mkdir(exist_ok=True)
z=json.load(open(R/'comparison.json'));D={(q['id'],q['interval']):q for q in z['records']}
main=['bg','richard','field','radial','m5_cap2.0','general_verified_cap0.1','general_cap0.2','general_cap0.3','general_verified_cap2.0','height_R0.3_cap0.1']
short={'bg':'BG only','richard':'Richard effective','field':'Uniform field','radial':'Uniform + radial','m5_cap2.0':'5-mode shared, relaxed','general_verified_cap0.1':'General stationary, ≤10%','general_cap0.2':'General stationary, ≤20%','general_cap0.3':'General stationary, ≤30%','general_verified_cap2.0':'General stationary, relaxed','height_R0.3_cap0.1':'Axial R≤300 nm, ≤10%'}
pages=[]
def page(title):pages.append([('h',title)])
def p(s):pages[-1].append(('p',s))
def eq(s):pages[-1].append(('eq',s))
def fig(name,width=6.8,height=None):pages[-1].append(('fig',name,width,height))
def table(headers,rows,widths=None):pages[-1].append(('table',headers,rows,widths))
def fnum(x):return '—' if x is None else f'{x:.1f}'
page('APD background and interference\nFit comparison for discussion with Richard')
p('Revised 6 October 2026 • Daping / Richard • Analysis of the supplied 40 s APD recording. This is a comparison of saved fits and sensitivity tests, not a calibration of absolute orientation. No new optimization was performed to prepare this report.')
p('<b>Main result.</b> The inferred background and polar-angle range depend strongly on the allowed field. The earlier ~55% background solution is not uniquely required: a broader stationary-field model reaches similar agreement at 20–30%, and the tested axial-motion model reaches similar agreement at 10%. These explanations cannot yet be distinguished with this recording alone.')
p('<b>What to discuss first.</b> Figures 1–2 show what the models actually fit; Figure 3 compares background makeup and angle ranges; Figure 5 shows how large background can coexist with the measured intensity through signed interference. The rest records the assumptions, controls and limits.')
rows=[]
for id in ['bg','richard','m5_cap2.0','general_cap0.3','general_verified_cap2.0','height_R0.3_cap0.1']:
 a,b=D[id,'30s'],D[id,'10s'];rows.append([short[id],f'{a["bg_max"]:.1f} / {b["bg_max"]:.1f}',f'{a["scores"]["test"]["xy_rms"]:.4f} / {b["scores"]["test"]["xy_rms"]:.4f}',f'{a["theta_min"]:.1f}–{a["theta_max"]:.1f}',f'{b["theta_min"]:.1f}–{b["theta_max"]:.1f}'])
table(['Model','BG / Tmax (%)\n30 s / 10 s','Held-out XY RMS\n30 s / 10 s','θ range (°)\n30–31 s','θ range (°)\n10–11 s'],rows,[155,82,108,70,70])
p('Direct inversion without background gives θ=27.4–57.7° at 30–31 s and 26.9–57.0° at 10–11 s. These are descriptive baselines, not known angles. All θ ranges are folded to 0–90° and sampled at the 128 fitted bins; none is a confidence interval.')
p('<b>Three conclusions supported by these calculations:</b> (i) simply changing the measured hole outline has very little effect in the tested pupil-mask model; (ii) the restricted hole-edge field tested here is insufficient; (iii) stationary-field structure and source position can trade against background power and recovered angles. A good cone fit alone cannot resolve that tradeoff.')
p('<b>Stopping point.</b> Retain NA<sub>in</sub>=0.38 and the current correction matrix. Do not choose a final background subtraction or publish absolute θ from the best residual alone. Decide which physical constraint or perturbation is worth measuring next.')
page('1. What the plotted averages mean')
p('The plots compare repeated traversals of the same measured anisotropy trajectory, taken from two windows of the 40 s APD recording. Each traversal retains its acquisition order, including the two branches. Channel intensities are averaged first; X and Y are calculated afterwards.')
table(['Plot label','What it contains','How it is used'],[['Training average','One alternating set of complete traversals: cycles 1,3,5,…','The model parameters are chosen using this average.'],['Held-out average','The other set: cycles 2,4,6,…','The fitted prediction is compared with this average without refitting.'],['Fitted prediction','The four-channel forward model evaluated along its fitted cone','The same prediction is shown against both averages.']],[105,230,150])
table(['Window','Training traversals','Held-out traversals'],[['30–31 s','152','152'],['10–11 s','109','108']],[165,160,160])
p('Both averages estimate the same trajectory within the same second. Agreement between them shows how reproducible the averaged curve is. A prediction that agrees with both is less likely to be following fluctuations peculiar to just one set. However, both sets share calibration, processing and any persistent distortion: this does not establish correct background or correct angles, and is not prediction of a new recording.')
p('The broad two-pass shape repeats; the 10–11 s window is more internally consistent than 30–31 s. All anisotropy axes remain −1 to +1. The plotted convention is:')
eq(r'X=\frac{I_0-I_{90}}{I_0+I_{90}},\qquad Y=\frac{I_{45}-I_{135}}{I_{45}+I_{135}}')
p('Fixed optical baseline: current gains and inverse T matrix, NA 0.38–1.30, water/oil indices 1.33/1.51, corrected Fresnel and BFP mapping. Detailed preprocessing and optimization records remain in the accompanying scripts; they are not needed to read these plots.')
page('2A. Common signal and power notation')
p('Here T means the sum of all four corrected channel intensities. S is rod-only intensity, B is background-only intensity, and C is the signed interference contribution. For every family:')
eq(r'I_j(\chi)=S_j(\chi)+C_j(\chi)+B_j')
eq(r'T(\chi)=\sum_j I_j(\chi)=S(\chi)+C(\chi)+B')
p('The instantaneous unit rod direction is d=(sinθ cosφ, sinθ sinφ, cosθ), and χ labels position around the cone. In field equations θ describes that direction over 0–180°; the reported angle plots fold it to 0–90°. This distinction preserves the sign of d_z when evaluating interference.')
eq(r'E_{s,j}(u,\chi)=a_0(d_x+i d_y)\sum_{l=x,y,z}d_l F_{jl}(u)')
p('F is the fixed collected dipole field, including the optics. The one constant a₀ per interval sets the overall rod brightness. The factor d_x+i d_y=sinθ exp(iφ) supplies the circular-excitation amplitude AND phase. It is not a freely chosen brightness at each point.')
eq(r'S_j=a_0^2\sin^2\theta\,q_j(\theta,\phi)')
eq(r'q_j=A+B_F\sin^2\theta+C_F\sin^2\theta\cos[2(\phi-\psi_j)]')
eq(r'S(\theta)=4a_0^2\sin^2\theta\,(A+B_F\sin^2\theta)')
p('For the annulus, ψ_j=0°,90°,45°,135°; A=0.89080598, B_F=0.10919402, C_F=0.92778057. Subscripts F distinguish the optical coefficients from background B and interference C. The shared-field code evaluates the corresponding pupil integrals numerically. Individual signal channels vary with θ and φ; their ideal annular sum varies with θ only.')
p('<b>Constant versus varying:</b> a₀ and the optical response are constant within an interval. θ and φ vary along the cone. In all stationary-background families below, B_j and their sum B are constant; it is C_j and C that may change with orientation. A negative C reduces measured power without making S or B negative.')
page('2B. Additive background and Richard’s effective model')
p('<b>1. Background only.</b> There is no modeled interference. Each channel has a constant additive intensity:')
eq(r'I_j=S_j+B_j,\qquad C_j=0')
eq(r'T=4a_0^2\sin^2\theta(A+B_F\sin^2\theta)+B')
p('The background is represented by a positive Stokes intensity pattern. This allows a constant polarization imbalance without making any channel negative:')
eq(r'B_j=\frac{B}{4}\{1+p\cos[2(\psi_j-\alpha)]\},\qquad 0\leq p\leq1')
p('B, polarization degree p and polarization angle α are constant. Total T varies only through the rod signal S(θ). This is the additive family used in the current forward-model comparison, not the exact historical shape-only optimization routine.')
p('<b>2. Richard effective intensity.</b> The implemented comparison adds one constant real overlap coefficient c_j per channel:')
eq(r'I_j=S_j+B_j+2c_j\sqrt{S_j B_j},\qquad -1\leq c_j\leq1')
eq(r'T=S(\theta)+B+2\sum_j c_j\sqrt{B_j S_j(\theta,\phi)}')
p('B_j and c_j remain constant. The interference varies because √S_j changes with orientation. Its sign in a given channel is fixed by c_j, although the total cross term may change sign as the differently weighted channels vary. Total intensity can acquire φ dependence even though the ideal rod-only total has none.')
p('This is equivalent to a constant coefficient times √S_j, with that coefficient written as 2c_j√B_j to enforce the scalar interference bound. The coefficient combines phase and spatial overlap; it is not necessarily cos of one physical phase. The fit does not separately identify a coherent and noninterfering share of B, nor guarantee that one common stationary field realizes all four coefficients.')
page('2C. Uniform field and uniform-plus-radial field')
p('<b>3. Uniform transverse background.</b> Let b_x and b_y be constant complex field amplitudes in normalized uniform pupil modes. Let U be the total extra noninterfering background. Then:')
eq(r'T=S(\theta)+2(|b_x|^2+|b_y|^2)+U+C_{\mathrm{uni}}(\theta,\phi)')
eq(r'C_{\mathrm{uni}}=4a_0\Gamma\sin^2\theta\,\mathrm{Re}\{e^{i\phi}(b_x^*\cos\phi+b_y^*\sin\phi)\}')
p('Γ is the fixed overlap between the collected transverse dipole field and a normalized uniform pupil field, computed from the annulus/Fresnel/BFP optics. It is not an extra adjustable overlap at each point. Both complex b coefficients and U are constant. The cross term changes magnitude and can change sign through θ and φ, even though the background field itself stays fixed.')
p('<b>4. Uniform plus radial background.</b> Add the normalized collected field pattern of a z-directed dipole, with constant complex amplitude b_r:')
eq(r'T=S(\theta)+2(|b_x|^2+|b_y|^2+|b_r|^2)+U+C_{\mathrm{uni}}+C_{\mathrm{rad}}')
eq(r'C_{\mathrm{rad}}=8a_0\kappa\sin\theta\cos\theta\,\mathrm{Re}(e^{i\phi}b_r^*),\qquad \kappa=\sqrt{A/2}')
p('The radial mode adds a different orientation dependence: sinθ cosθ times a φ-dependent phase projection. Its intensity is still constant. In these normalized, mutually orthogonal pupil modes, summed coherent background power is twice the sum of squared field amplitudes. Cross terms between background modes need not vanish in every individual analyzer channel; they are retained in the channel calculation.')
p('U is distributed across channels by the same positive Stokes form as above. “Noninterfering” means no cross term with the modeled rod signal: it can include incoherent light or a coherent field with no relevant spatial overlap. These expressions describe a background family, not a claim that a uniform or radial pattern came from a particular optical surface.')
p('The separate-interval three-mode fits allow their field coefficients to change between windows. The shared three-mode test fixes the same coefficients and U across both windows, while each window retains its own cone and a₀. A budget cap changes the allowed power, not these equations.')
page('2D. Five modes, general stationary field and hole edge')
p('These families have the same power equation. They differ in which fixed spatial/polarization patterns g_m are available to form the background:')
eq(r'E_{b,j}(u)=\sum_{m=1}^{M}b_m g_{jm}(u)')
eq(r'H^{\Sigma}_{lm}=\sum_j\int F_{jl}(u)g_{jm}^*(u)\,du')
eq(r'T=S(\theta)+2\sum_{m=1}^{M}|b_m|^2+U+C_M(\theta,\phi)')
eq(r'C_M=2a_0\,\mathrm{Re}\!\left[(d_x+i d_y)\sum_{l,m}d_l H^{\Sigma}_{lm}b_m^*\right]')
p('The modes are orthonormal before analyzer projection, giving the factor 2 in total background power for this four-analyzer convention. H is a precomputed optical overlap, constant for a stationary source. The b_m and U are constant. Thus orientation changes C_M through the displayed products of direction components; no per-point overlap is fitted.')
table(['Family','Allowed background pattern','Final total-intensity form'],[['5-mode shared','Uniform x/y plus orthogonalized collected z-, x- and y-dipole patterns','T=S + 2Σ₁⁵|bₘ|² + U + C₅'],['General stationary','Each transverse polarization uses the five-dimensional scalar span of the collected dipole components (10 modes)','T=S + 2Σ₁¹⁰|bₘ|² + U + C₁₀'],['Hole-edge surrogate','Two transverse polarizations × three edge-localized angular patterns (6 modes)','T=S + 2Σ₁⁶|bₘ|² + C₆; U=0']],[100,265,120])
p('The edge patterns use exp[−(ρ−0.38)²/(2w²)] times 1, cosϕ_u or sinϕ_u, then normalization/orthogonalization. Here ϕ_u is pupil azimuth, not rod azimuth φ. Width w is fixed for each tested case. The six complex amplitudes are shared across both windows, as are the five-/ten-mode coefficients in the corresponding shared fits.')
p('For a stationary field, C_M is quadratic in d and has mechanical-phase harmonics only up to order two on the cone; S is quartic and can reach order four. Increasing the mode family changes the allowed quadratic coefficients, not the permission to choose an arbitrary C at every bin. The ten-mode completeness statement assumes a fixed-position ideal dipole, common pupil and full collection.')
page('2E. Axial motion, mask and deliberately flexible controls')
p('<b>5. Bounded axial motion with a stationary background.</b> The ten background coefficients stay fixed. The rod position changes its phase relative to that field, so the optical overlap now depends on z:')
eq(r'z(\chi)=R\sin\theta_{\mathrm{axis}}\cos(\chi+\delta),\quad f(z)=[1+(z/z_R)^2]^{-1/2}')
eq(r'H^{\Sigma}_{lm}(z)=\sum_j\int F_{jl}g_{jm}^*e^{ikz[n_w+\sqrt{n_w^2-\rho^2}]}\,du')
eq(r'T=f(z)^2S(\theta)+2\sum_{m=1}^{10}|b_m|^2+U+C_{10}(\theta,\phi,z)')
eq(r'C_{10}=2a_0 f(z)\,\mathrm{Re}\!\left[e^{-i\arctan(z/z_R)}(d_x+i d_y)\sum_{l,m}d_l H^{\Sigma}_{lm}(z)b_m^*\right]')
p('R and δ are constants per interval; z_R=2 µm and k=2π/(633 nm) are fixed in these tests. Background intensity is still constant. θ, φ and z vary together along the trajectory. The signal acquires f²; interference acquires f and a deterministically changing propagation/Gouy phase. There is no independent focus multiplier per point. The separate finite-aperture simulation changes collection itself and was not included in this fitted power equation.')
p('<b>Measured-mask variant.</b> The mask changes the fixed signal response and overlaps, not the basic power balance:')
eq(r'T=a_0^2|d_x+i d_y|^2\,\mathbf{d}^{T}Q^{\Sigma}_{\mathrm{mask}}\mathbf{d}+C_{\mathrm{mask}}(\mathbf{d})+B')
p('Q and H are recomputed through the measured aperture, and B remains constant. The asymmetric mask may introduce φ dependence even into rod-only total power. The mask fits used the three-mode background family; an amplitude mask does not model light scattered from the hole wall.')
p('<b>Arbitrary-overlap feasibility control, not a stationary-field fit:</b>')
eq(r'T=S+B+2\sum_j\rho_j(\chi)\sqrt{S_j(\chi)B_j},\qquad |\rho_j(\chi)|\leq1')
p('Here B_j stay constant but every ρ_j can vary independently at every point. That freedom can reproduce the data exactly without demonstrating one realizable stationary field. The earlier free-brightness models additionally replaced the prescribed a₀ sinθ by an independently recovered amplitude a(χ):')
eq(r'T=a(\chi)^2Q(\theta,\phi)+a(\chi)L(\theta,\phi)+B')
p('Q is the summed collection response for unit excitation amplitude; L is the cross-term coefficient for the selected background family. These are historical relaxations, not the main fixed-excitation results.')
page('3. What the older model families can reproduce')
fig('01_model_fits',6.65,7.15)
p('Figure 1. Points are held-out cycle averages; grey is the training average; red is the saved prediction. The additive and Richard-effective fits are separate per interval. The five-mode field uses the same absolute background coefficients in both intervals. All three use constant a₀ with circular-excitation dependence. Richard’s model need not win on XY RMS because that is not the channel fitting objective.')
page('4. Comparable fits with smaller background budgets')
fig('02_low_background_fits',6.65,7.15)
p('Figure 2. General stationary field at 10% and 30% caps, and the same ten-mode family with bounded axial motion at 10%. The last row is a sensitivity scenario, not evidence that the fitted motion actually occurs. These curves show why the earlier large-background estimate cannot be treated as a data-only requirement.')
page('5. Background makeup and the angle consequences')
fig('03_background_and_angles',6.85,3.35)
p('Figure 3. Background is normalized to the maximum training total, T=∑I<sub>j</sub>. The decomposition bars show 30–31 s; exact values for both windows are tabulated in the appendix. Grey angle bands show direct inversion. Changing the field family changes θ substantially, not just the residual.')
p('The five-mode shared relaxed fit has B/Tmax=54.72% / 54.43% (30 s / 10 s). Its coherent share is 91.82%, giving coherent power 50.25% / 49.97% of Tmax and noninterfering power 4.47% / 4.45%. In the general stationary and axial-motion selected fits, the residual Stokes component is effectively zero: essentially all fitted background is in the represented coherent field.')
p('This split is conditional on the chosen mode space. “Noninterfering” can include light temporally incoherent with the rod, or coherent light spatially orthogonal to every signal mode being tested. Conversely, the represented coherent field can contain components with weak overlap at some orientations. Neither split assigns a percentage to coverslip, cell, objective or hole wall.')
p('Historical “about 90–95% background” referred to the <b>minimum</b> measured total in some earlier fits. It does not mean 90–95% of the maximum. Both denominators are reported in the appendix. Nor does B≈T at a dim point imply S≈0 when C is negative: T=S+C+B, so destructive interference can conceal substantial rod-only power.')
p('The unrestricted-overlap construction is deliberately excluded from these fitted-angle bars: it can reproduce the training curve exactly at a 10% background budget, but it chooses overlap independently at each point. That demonstrates feasibility under weak bounds, not a valid stationary-field correction or validated angle recovery.')
page('6. Power decomposition: why large background can survive')
fig('05_power_decomposition',6.85,5.5)
p('Figure 5. All terms use the same Tmax normalization within each window. The rod-alone term obeys the one-a₀ excitation and collection law. The constant background and signed interference can compensate over substantial parts of the trajectory; the amount is not confined by an assumed positive “interference fraction.” Exact per-bin values are in comparison.json.')
p('This checks the proposed intensity consistency condition: the recovered θ and the rod-only intensity use the same forward model and scale. Passing that internal consistency check does not establish the true θ, because cone geometry, field overlap and background were jointly chosen. The general 30% and moving-source 10% examples redistribute the same observed power differently.')
page('7. Angle variation along the trajectory')
fig('06_angle_traces',6.85,4.75)
p('Figure 6. θ is folded to 0–90°. Azimuth differences are wrapped modulo 180°, with plot breaks at wrap discontinuities. These are comparisons at matched ordered bins, not errors relative to ground truth. Full θ and φ arrays, wrapped φ spans, and RMS differences from direct inversion are supplied in JSON/CSV.')
p('The same measured curve admits fitted θ ranges extending almost to the equator in relaxed models, roughly 31–65° at the 30% general stationary cap, or roughly 35–55° in the 300 nm axial-motion scenario. Near-90° solutions are especially sensitive to intensity/model errors. The evidence supports substantial angle ambiguity; it does not support choosing the narrowest or widest interval as more physical.')
p('Do not interpret a wrapped φ minimum–maximum as a continuous rotation amplitude. The modulo-180° trace and branch-aware difference are the appropriate comparisons. A cone imposes continuity and a branch convention but does not add an independently known orientation.')
page('8. Stationary-field completeness and the budget tradeoff')
fig('04_fit_tradeoffs',6.85,2.8)
p('Figure 4. Left: selected general stationary fits versus their background budget, with the previous relaxed five-mode fit marked ×. Right: adding bounded axial motion at 10% background. The 10–11 s held-out residual need not decrease monotonically because optimization targets training channels, not held-out XY.')
p('For a fixed-position ideal dipole, form the scalar span V of the x/y Cartesian components of its three collected dipole fields. Its numerical rank is five. Two transverse background components in V require ten complex coefficients. The Cartesian remainder is orthogonal to V: it has no interference with any signal orientation and no integrated cross term with the represented field. Its observable residual intensity is captured by three positive Stokes parameters.')
p('This supplies a complete observation representation for a <b>stationary common-pupil field under the fixed-position, ideal-dipole, full-collection assumptions</b>. Ten is a sufficient mode count, not a count of physical scattering objects or uniquely identifiable coefficients. Basis-dependent modal powers should not be labeled as specific sources. Adding further stationary pupil modes cannot improve this ideal fixed-position model beyond that representation.')
p('With the circular-excitation law, the stationary rod–background cross term is quadratic in the direction d: on a fixed cone it contains mechanical-phase harmonics only through order two. Rod-only intensity is quartic and can contain order four. This restriction is why an arbitrary overlap at every bin is much more permissive than a stationary field.')
p('The fit remains nonconvex in cone, phase and field parameters. Reported minima are feasible examples from the saved multistart searches, not proofs that a lower budget is impossible. Independent post-analyzer backgrounds, calibration errors, non-dipolar response, changing excitation and moving-source effects fall outside the completeness claim. Alternate-cycle agreement alone does not resolve them.')
page('9. Hole shape, localized scattering and bounded height')
fig('07_edge_and_focus',6.85,2.8)
p('Figure 7. Left: six complex edge-field coefficients, shared across windows, with Gaussian radial envelope exp[−(ρ−0.38)²/(2w²)] multiplied by 1, cosϕ, sinϕ in each transverse polarization. Widths 0.03, 0.10 and 0.25 NA were tested. Narrow fields fit poorly; broader fields improve but become less specifically localized to the edge. The narrow 10% selected run reached its evaluation limit; it is not a converged exclusion result.')
p('This is an effective low-order field family, not a Maxwell calculation of the drilled wall. Wall topography, coating, mirror tilt and distance from the true BFP are not known. It cannot rule out all hole scattering. Separately, using the thesis measured amplitude mask at four rotations changed held-out XY RMS by less than about 0.0002 in the three-mode comparisons. That tests pupil transmission shape, not edge scattering or out-of-BFP propagation.')
p('Axial-motion model: z(χ)=R sin(θ<sub>axis</sub>)cos(χ+δ), with R and δ independent per window. R is not the rod-orientation cone radius or assumed rod length. Caps 50, 150 and 300 nm are sensitivity scenarios, not biological measurements. The 300 nm fits use R≈228 / 300 nm, but axial amplitudes only 40.7 / 43.0 nm (about 81 / 86 nm peak-to-peak). Boundary-hitting radii warn against treating them as recovered hook geometry.')
p('The rod field acquires axial phase exp[i(2π/λ)z(n<sub>w</sub>+√(n<sub>w</sub>²−ρ²))], λ=633 nm, including incident and collected propagation. A Gaussian illumination with 2 µm Rayleigh range and its Gouy phase is assumed. At the selected amplitudes its intensity change is below 0.05%; phase-dependent interference, rather than brightness loss, drives most of this fitted improvement.')
p('Only axial motion is modeled; lateral motion, uncertain mean height, interface-dependent rod response and finite APD acceptance are not fitted jointly. Once the source moves, the ten-mode background is a restricted family again; fixed-position completeness no longer applies. The separate pure-defocus aperture simulation on the right tests ±300 nm and several orientations, without background. Full collection conserves power; clipping can alter θ. These aperture radii are not calibrated APD sizes.')
page('10. Physical interpretation and the useful stopping line')
p('<b>Background categories from Richard’s correspondence.</b> External optical reflections/scattering can carry arbitrary phase and spatial structure. Ideal coverslip reflection should leave through the central hole. Any reflected light that remains after bouncing or scattering need not retain a simple circular-polarization phase relation. Cell scattering can come from different heights and different polarization components; nearby rods add sample-associated fields. The APD integrates their overlap, so equal source position for the rod does not imply equal interference phase in all channels.')
p('Multiple mutually coherent stationary backgrounds add as fields before squaring. If they lie in the same allowed mode space, fitting two copies does not add a new observation model; it only makes the decomposition redundant. Mutually incoherent contributions add intensities. More sources become observationally distinct when they add new spatial/polarization modes, different coherence, motion or a response to an experimental perturbation. Their source locations cannot generally be inferred from one steady four-channel trajectory.')
p('Fresnel transmission, BFP mapping and the hole act on fields along the optical path; a source should enter where it is generated. After propagation to one common detection plane, the background can be represented there for this effective model. Intensity subtraction is not generally a commuting substitute for coherent propagation. The present fits do not identify which optical surface produced the effective field.')
p('<b>What the existing data can do:</b> reject or expose poor versions of restricted model families; show budget–angle sensitivity; check shared-field consistency across windows; quantify registration and chronological variation; test whether physically linked terms reproduce channels as well as the outline. It cannot establish unique absolute θ, actual coherent source fractions, APD acceptance, or true centre motion without additional constraints.')
p('<b>Priority discussion with Richard:</b> which perturbation best separates stationary optical field from sample-associated field in this APD setup? An axial stage scan/modulation with four-channel acquisition is especially direct, provided rod orientation, position and exposure averaging are tracked. It can constrain a phase-dependent cross term; it is not automatically an exact background subtraction.')
p('Fixed-rod files can quantify temporal stability, channel covariance, spectra and response to laser-power steps. Flat intensity alone cannot bound a static coherent background. A nearly constant-θ rotating rod can test azimuth-dependent power and shared-field predictions, but its circular anisotropy provides less leverage for separating additive scale/background. These are useful controls, not substitutes for known-angle truth.')
p('Other discriminating measurements are dark electronic offsets, a current transmission/mixing calibration, effective APD acceptance/defocus scans, and centre-position tracking. The sensible next step is to pick one or two constraints, not add another unconstrained per-bin term. A synthetic injection/recovery study can characterize an estimator under assumed models, but would not decide which model describes this old sample.')
page('11. Source limits, verification and reproducible record')
p('<b>Bearing paper and SI.</b> The supplied manuscript and supplementary-information (1).pdf describe deliberately selected near-axis rods with overlapping loops, fewer than 10% of candidates. Do not transfer that geometry or biological conclusions to ordinary BFM samples. SI S1 documents the drilled mirror near an inaccessible BFP, hole-wall scattering, polarization/alignment issues and APD alignment. SI p5 transmission/extinction values are historical, not a current calibration. Fixed-rod noise measurements and manual azimuth tests establish precision/φ behavior, not exact θ or a coherent-background bound.')
p('<b>Polarcam draft and SI.</b> Polarcam Manuscript V2 2026-09-08 RB.docx, SI outline.docx and RB Figs V2.docx were reviewed separately. Their camera background values, compensation optics, exposure/modulation behavior and some placeholder numbers are not transferred to this APD data. Axial modulation and lateral minimum projection motivate discriminating tests but require their own assumptions. APD on the Polarcam setup is not the bearing-paper APD setup.')
p('<b>Thesis and optical sources.</b> Thesis sections 1.5–1.8 and its measured mirror mask supplied optical context; mask provenance is thesis commit 7576b50, analysis/image_mirror_binary.tif. The active baseline remains NA<sub>in</sub>=0.38. Original notebook fitting history was recovered from repository import 9b3604e. The model conventions and dated source notes accompany the scripts. General pupil-phase context: https://pmc.ncbi.nlm.nih.gov/articles/PMC5810908/ ; the present calculation uses water-side propagation with the existing water–oil transmission.')
p('<b>Numerical checks.</b> This report reconstructs and checks 46 model/window predictions against saved residuals and checks S+C+B consistency. Earlier checks reproduce all 15 saved edge/general/height comparisons; arbitrary fixed-field overlap projection error is about 2.9×10⁻¹⁶; pupil refinement changes channels by about 1.2×10⁻⁴ for general fields and 7.2×10⁻⁴ for the narrow edge family; checked axial interpolation error is about 5.3×10⁻⁷. These are numerical checks, not experimental accuracy bounds.')
p('The best narrow-edge 10% run has success=false at 600 function evaluations. Other selected runs in the principal table meet their optimizer termination criteria. Several caps/radius constraints are active, and multistart optimization does not certify global minima. No information criterion or formal confidence interval is claimed: correlated bins, estimated registration and unknown error covariance make naive parameter-count significance unreliable.')
p('<b>Reproduction.</b> comparison.csv is a compact numerical table; comparison.json includes per-bin predictions, signal/interference/background decomposition, training/test arrays and angles. build_data.py reconstructs results; plot_figures.py draws seven PNG/SVG figure sets; build_report.py creates this PDF and editable Word document. The study archive preserves earlier scripts, fit parameters, cached averages, optical inputs and dated notes. Replotting cached results does not require the large raw TDMS. Reprocessing and refitting require their documented data paths and dependencies.')
p('No future Codex handoff script is included yet. The code and dated record are the current milestone; the report is intended to settle the scientific next step before extending the model.')
# Full selected-model tables. Exclude duplicated old/verified general minima from display, retain in machine data.
ids=[a['id'] for a in z['records'] if a['interval']=='30s' and a['id'] not in ['general_cap0.1','general_cap2.0']]
for ki,win in [('30s','30–31 s'),('10s','10–11 s')]:
 page(f'Appendix A · {win}: fit quality and angles')
 rows=[]
 for id in ids:
  q=D[id,ki];rows.append([q['label'],f'{q["scores"]["test"]["xy_rms"]:.4f}',f'{100*q["scores"]["test"]["channel_rms"]:.2f}',f'{100*q["scores"]["test"]["total_rms"]:.2f}',f'{q["theta_min"]:.1f}–{q["theta_max"]:.1f}',f'{q["phi_rms_vs_direct"]:.1f}',str(q['n_parameters'])+(' S' if q['scope']=='shared field' else ' L')+(' *' if not q['success'] else '')])
 table(['Model','XY RMS','Channel\nRMS %','Total\nRMS %','θ span (°)','φ RMS vs\ndirect (°)','Param.'],rows,[171,48,50,48,67,59,42])
 p('Channel and total RMS use residual divided by observed total at each bin; channel RMS is not relative to each channel individually. θ span covers the sampled fitted orientations. φ differences use modulo 180°. S = total parameters in the joint two-window fit; L = parameters in this separate single-window fit. * = selected run reached its evaluation limit. Numeric termination does not establish global optimality.')
 p('Direct no-background inversion is not a fitted cone and has no fitted prediction score here. All historical and verified duplicates are retained in comparison.json/CSV. Early free-brightness and arbitrary-overlap constructions are not comparable predictive models and are intentionally kept outside this table.')
 page(f'Appendix B · {win}: background composition')
 rows=[]
 for id in ids:
  q=D[id,ki];rows.append([q['label'],f'{q["bg_max"]:.2f}',f'{q["bg_min"]:.2f}',fnum(q['coherent_max']),fnum(q['residual_max']),'/'.join(f'{b:.1f}' for b in q['bg_channels_max'])])
 table(['Model','All BG\n% Tmax','All BG\n% Tmin','Coherent\n% Tmax','Residual\n% Tmax','Channel BG % Tmax\n0 / 90 / 45 / 135'],rows,[172,50,50,57,55,101])
 p('Tmax and Tmin are from the training total within this window. Channel entries sum to all BG up to rounding. Coherent/residual components are nonnegative fitted powers; interference is signed and reported separately in Figure 5 and comparison.json. A dash means the effective model does not identify that split. Near-zero residual components are rounded to 0.0%; this is a model-dependent boundary solution, not proof that all real background is coherent.')
 p('Shared-field coefficients are identical in absolute units between windows. Their percentages differ because the observed maxima/minima differ. The background is constant along each fitted trajectory; it is the signal–background cross term that varies with orientation and, in the height scenario, position.')
# Write PDF and DOCX with common structured content.
fontdir=Path('/usr/share/fonts/truetype/dejavu');pdfmetrics.registerFont(TTFont('DejaVu',str(fontdir/'DejaVuSans.ttf')));pdfmetrics.registerFont(TTFont('DejaVu-Bold',str(fontdir/'DejaVuSans-Bold.ttf')));pdfmetrics.registerFontFamily('DejaVu',normal='DejaVu',bold='DejaVu-Bold',italic='DejaVu',boldItalic='DejaVu-Bold')
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyCustom',fontName='DejaVu',fontSize=9,leading=13,spaceAfter=9,textColor=colors.HexColor('#253343')));styles.add(ParagraphStyle(name='HeadCustom',fontName='DejaVu-Bold',fontSize=18,leading=23,spaceAfter=17,textColor=colors.HexColor('#123d54')));styles.add(ParagraphStyle(name='CellCustom',fontName='DejaVu',fontSize=7.2,leading=10));styles.add(ParagraphStyle(name='CellHeadCustom',fontName='DejaVu-Bold',fontSize=7.2,leading=10,textColor=colors.white))
story=[];doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7);sec.page_width=Inches(8.5);sec.page_height=Inches(11)
doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
plain=lambda s:html.unescape(re.sub('<[^>]+>','',s))
for ii,blocks in enumerate(pages):
 if ii:story.append(PageBreak());doc.add_page_break()
 for block in blocks:
  typ=block[0]
  if typ in ['p','h']:
   text=block[1];story.append(Paragraph(text.replace('\n','<br/>'),styles['HeadCustom' if typ=='h' else 'BodyCustom']));doc.add_paragraph(plain(text),style='Heading 1' if typ=='h' else None)
  elif typ=='eq':
   import matplotlib
   matplotlib.use('Agg')
   import matplotlib.pyplot as plt
   from PIL import Image as PILImage
   eqdir=OUT/'equations';eqdir.mkdir(exist_ok=True)
   import hashlib
   path=eqdir/(hashlib.sha256(block[1].encode()).hexdigest()[:16]+'.png')
   ef=plt.figure(figsize=(1,.5));ef.text(0,.1,'$'+block[1]+'$',fontsize=14)
   ef.savefig(path,dpi=240,bbox_inches='tight',pad_inches=.07,facecolor='white');plt.close(ef)
   iw,ih=PILImage.open(path).size;scale=min(72/240,510/iw);ww,hh=iw*scale,ih*scale
   story.extend([Image(str(path),width=ww,height=hh,hAlign='LEFT'),Spacer(1,9)])
   doc.add_picture(str(path),width=Inches(ww/72))
  elif typ=='fig':
   _,name,w,h=block;path=R/'figures'/f'{name}.png'
   from PIL import Image as PILImage
   iw,ih=PILImage.open(path).size;scale=min(w*72/iw,(h*72/ih if h else 100));ww,hh=iw*scale,ih*scale;story.append(Image(str(path),width=ww,height=hh));story.append(Spacer(1,9));doc.add_picture(str(path),width=Inches(ww/72))
  else:
   _,heads,rows,widths=block;cells=[[Paragraph(html.escape(str(c)).replace('\n','<br/>'),styles['CellHeadCustom' if i==0 else 'CellCustom']) for c in row] for i,row in enumerate([heads]+rows)]
   t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#194c60')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#f0f5f7'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#bbc9cf'))]));story.extend([t,Spacer(1,12)])
   dt=doc.add_table(rows=1,cols=len(heads));dt.style='Light Shading Accent 1'
   for c,s in zip(dt.rows[0].cells,heads):c.text=s
   for row in rows:
    for c,s in zip(dt.add_row().cells,row):c.text=str(s)
   for row in dt.rows:
    for c in row.cells:
     for pp in c.paragraphs:
      for run in pp.runs:run.font.size=Pt(8)
   doc.add_paragraph()
def footer(canvas,document):
 canvas.setFont('DejaVu',7);canvas.setFillColor(colors.HexColor('#607381'));canvas.drawString(48,25,'APD background / interference • Discussion record • 6 October 2026');canvas.drawRightString(564,25,str(document.page))
SimpleDocTemplate(str(OUT/'APD_background_discussion_2026-10-06.pdf'),pagesize=(612,792),leftMargin=43,rightMargin=43,topMargin=38,bottomMargin=43,title='APD background and interference: fit comparison',author='Analysis prepared for Daping and Richard').build(story,onFirstPage=footer,onLaterPages=footer)
doc.save(OUT/'APD_background_discussion_2026-10-06.docx')
(R/'report_content.json').write_text(json.dumps(pages,indent=2))
print('Created report with',len(pages),'planned pages.')
